import hashlib
import html
import json
import pickle
import re
import time
import urllib.error
import urllib.request
from pathlib import Path

import docx
import faiss
import google.generativeai as genai
import numpy as np
import pandas as pd
import pypdf
import streamlit as st
from rank_bm25 import BM25Okapi
from sentence_transformers import CrossEncoder, SentenceTransformer


# ── Constants ──────────────────────────────────────────────────────────────
APP_DIR = Path(__file__).resolve().parent
CACHE_DIR = APP_DIR / ".rag_cache"
CHAT_DIR = APP_DIR / ".chat_history"
EMBEDDING_MODEL_NAME = "all-MiniLM-L6-v2"
RERANKER_MODEL_NAME = "cross-encoder/ms-marco-MiniLM-L-6-v2"
GEMINI_MODEL_NAME = "gemini-2.5-flash"
GROQ_MODEL_NAME = "llama-3.3-70b-versatile"
GROQ_ENDPOINT = "https://api.groq.com/openai/v1/chat/completions"
# Cloudflare (in front of Groq) blocks the default Python-urllib User-Agent (error 1010).
HTTP_USER_AGENT = "Mozilla/5.0 (compatible; RAGAnalyzer/1.0; +https://streamlit.io)"
PARSER_VERSION = "v2-page-sheet-citations"
TOP_K_RETRIEVAL = 10
TOP_K_CONTEXT = 4
MAX_FILE_SIZE_MB = 50
MAX_TOTAL_UPLOAD_MB = 150
MAX_CHUNKS_PER_FILE = 4000
LLM_MAX_RETRIES = 2
LLM_RETRY_BASE_DELAY_SECONDS = 2


# ── Page config & CSS ──────────────────────────────────────────────────────
st.set_page_config(
    page_title="RAG Analyzer",
    page_icon="🔍",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&display=swap');

html, body, [class*="css"] {
    font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif !important;
}

/* ── Base ── */
.stApp { background: linear-gradient(180deg, #ffffff 0%, #fafafd 100%); }
#MainMenu { visibility: hidden; }
footer { visibility: hidden; }
header[data-testid="stHeader"] { background: transparent; }
.block-container { padding-top: 2.2rem !important; }

/* ── Sidebar ── */
[data-testid="stSidebar"] {
    background-color: #f5f5f7 !important;
    border-right: 1px solid #e5e5e5 !important;
}
[data-testid="stSidebar"] .stMarkdown p {
    font-size: 0.82rem;
    color: #6e6e73;
}
[data-testid="stSidebar"] h3 {
    font-size: 0.62rem !important;
    font-weight: 700 !important;
    letter-spacing: 0.09em !important;
    text-transform: uppercase !important;
    color: #86868b !important;
    margin-top: 24px !important;
    margin-bottom: 8px !important;
    padding-bottom: 6px !important;
    border-bottom: 1px solid #e5e5e5 !important;
}

/* ── Hero banner ── */
.hero {
    background: linear-gradient(135deg, #1d1d1f 0%, #2b2b30 55%, #3a3a40 100%);
    border-radius: 22px;
    padding: 30px 34px;
    margin-bottom: 22px;
    color: #ffffff;
    box-shadow: 0 12px 40px rgba(0,0,0,0.18);
    position: relative;
    overflow: hidden;
}
.hero::after {
    content: "";
    position: absolute;
    top: -60px; right: -40px;
    width: 220px; height: 220px;
    background: radial-gradient(circle, rgba(0,113,227,0.35) 0%, rgba(0,113,227,0) 70%);
}
.hero-badge {
    display: inline-block;
    font-size: 0.64rem;
    font-weight: 700;
    letter-spacing: 0.12em;
    text-transform: uppercase;
    color: #ffffff;
    background: rgba(255,255,255,0.14);
    border: 1px solid rgba(255,255,255,0.18);
    padding: 5px 12px;
    border-radius: 980px;
    margin-bottom: 14px;
}
.hero-title {
    font-size: 2.1rem;
    font-weight: 800;
    letter-spacing: -0.045em;
    margin: 0 0 8px;
    line-height: 1.05;
}
.hero-desc {
    font-size: 0.95rem;
    color: #d2d2d7;
    line-height: 1.6;
    max-width: 640px;
    margin: 0;
    position: relative;
    z-index: 1;
}
.hero-desc strong { color: #ffffff; font-weight: 600; }

/* ── Metrics ── */
[data-testid="stMetric"] {
    background: #ffffff !important;
    border-radius: 14px !important;
    padding: 14px 18px !important;
    border: 1px solid #ececf0 !important;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04) !important;
    transition: transform 0.15s ease, box-shadow 0.15s ease !important;
}
[data-testid="stMetric"]:hover {
    transform: translateY(-2px) !important;
    box-shadow: 0 6px 18px rgba(0,0,0,0.08) !important;
}
[data-testid="stMetricLabel"] p {
    font-size: 0.62rem !important;
    font-weight: 700 !important;
    text-transform: uppercase !important;
    letter-spacing: 0.07em !important;
    color: #86868b !important;
}
[data-testid="stMetricValue"] {
    font-size: 1.05rem !important;
    font-weight: 600 !important;
    color: #1d1d1f !important;
}

/* ── Buttons ── */
.stButton > button {
    border-radius: 980px !important;
    font-size: 0.8rem !important;
    font-weight: 500 !important;
    border: 1px solid #d2d2d7 !important;
    background: #ffffff !important;
    color: #1d1d1f !important;
    padding: 6px 16px !important;
    transition: background 0.15s ease, border-color 0.15s ease !important;
    box-shadow: none !important;
}
.stButton > button:hover {
    background: #ebebeb !important;
    border-color: #c7c7cc !important;
    transform: none !important;
    box-shadow: none !important;
}

/* ── Chat input ── */
[data-testid="stChatInputTextArea"] {
    border-radius: 22px !important;
    border: 1px solid #d2d2d7 !important;
    background: #f5f5f7 !important;
    font-size: 0.9rem !important;
}
[data-testid="stChatInputTextArea"]:focus {
    border-color: #0071e3 !important;
    background: #ffffff !important;
    box-shadow: 0 0 0 3px rgba(0,113,227,0.12) !important;
    outline: none !important;
}

/* ── Chat messages ── */
[data-testid="stChatMessage"] {
    background: transparent !important;
    border: none !important;
    padding: 4px 0 !important;
}
/* Hide default avatars for a cleaner iMessage-style thread */
[data-testid="stChatMessage"] [data-testid="stChatMessageAvatarUser"],
[data-testid="stChatMessage"] [data-testid="stChatMessageAvatarAssistant"],
[data-testid="stChatMessage"] [data-testid="chatAvatarIcon-user"],
[data-testid="stChatMessage"] [data-testid="chatAvatarIcon-assistant"] {
    display: none !important;
}

/* User bubble — right aligned */
[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) {
    flex-direction: row-reverse !important;
}
[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) [data-testid="stChatMessageContent"],
[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) [data-testid="stMarkdownContainer"] {
    background: linear-gradient(135deg, #0071e3 0%, #0a84ff 100%) !important;
    color: #ffffff !important;
    border-radius: 20px 20px 6px 20px !important;
    padding: 11px 16px !important;
    max-width: 78% !important;
    box-shadow: 0 2px 10px rgba(0,113,227,0.22) !important;
}
[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) [data-testid="stMarkdownContainer"] * {
    color: #ffffff !important;
}

/* Assistant bubble — left */
[data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-assistant"]) [data-testid="stMarkdownContainer"] {
    background: #f2f2f5 !important;
    border: 1px solid #e9e9ee !important;
    border-radius: 6px 20px 20px 20px !important;
    padding: 12px 17px !important;
    max-width: 88% !important;
    color: #1d1d1f !important;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04) !important;
}

/* ── File uploader ── */
[data-testid="stFileUploader"] {
    background: #ffffff !important;
    border: 2px dashed #c7c7cc !important;
    border-radius: 12px !important;
    padding: 6px !important;
    transition: border-color 0.2s ease !important;
}
[data-testid="stFileUploader"]:hover,
[data-testid="stFileUploader"]:focus-within {
    border-color: #0071e3 !important;
}
[data-testid="stFileUploaderDropzone"] {
    background: #f5f5f7 !important;
    border-radius: 10px !important;
}

/* ── Expander ── */
[data-testid="stExpander"] summary {
    font-size: 0.82rem !important;
    font-weight: 500 !important;
    color: #1d1d1f !important;
}

/* ── Download button — black Apple pill ── */
[data-testid="stDownloadButton"] > button {
    border-radius: 980px !important;
    font-size: 0.8rem !important;
    font-weight: 600 !important;
    border: 1px solid #1d1d1f !important;
    background: #1d1d1f !important;
    color: #ffffff !important;
    padding: 7px 20px !important;
    transition: background 0.15s ease, transform 0.15s ease !important;
}
[data-testid="stDownloadButton"] > button:hover {
    background: #000000 !important;
    transform: translateY(-1px) !important;
}

/* ── Welcome / landing screens ── */
.welcome-wrap {
    display: flex;
    flex-direction: column;
    align-items: center;
    justify-content: center;
    min-height: 58vh;
    text-align: center;
    padding: 40px 24px;
}
.welcome-icon { font-size: 56px; margin-bottom: 20px; line-height: 1; }
.welcome-title {
    font-size: 2.6rem;
    font-weight: 800;
    color: #1d1d1f;
    letter-spacing: -0.04em;
    margin: 0 0 12px;
}
.welcome-sub {
    font-size: 1.05rem;
    color: #6e6e73;
    line-height: 1.6;
    max-width: 460px;
    margin: 0 auto 40px;
}
.feature-grid {
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 12px;
    max-width: 540px;
    width: 100%;
    text-align: left;
}
.feature-card {
    background: #ffffff;
    border: 1px solid #ececf0;
    border-radius: 16px;
    padding: 20px;
    box-shadow: 0 1px 3px rgba(0,0,0,0.04);
    transition: transform 0.15s ease, box-shadow 0.15s ease;
}
.feature-card:hover {
    transform: translateY(-3px);
    box-shadow: 0 10px 26px rgba(0,0,0,0.09);
}
.feature-card-title {
    font-size: 0.84rem;
    font-weight: 600;
    color: #1d1d1f;
    margin: 0 0 5px;
}
.feature-card-desc {
    font-size: 0.76rem;
    color: #6e6e73;
    line-height: 1.5;
    margin: 0;
}

/* ── Doc card ── */
.doc-card {
    background: #ffffff;
    border: 1px solid #e5e5e5;
    border-radius: 10px;
    padding: 12px 16px;
    margin-bottom: 8px;
    font-size: 0.83rem;
    color: #1d1d1f;
}

/* ── Mobile hint (shown only on small screens) ── */
.mobile-hint {
    display: none;
    background: #fff8e6;
    border: 1px solid #ffe08a;
    color: #7a5b00;
    border-radius: 12px;
    padding: 11px 15px;
    font-size: 0.82rem;
    line-height: 1.45;
    margin-bottom: 16px;
}

/* ── Central intake / search-bar composer ── */
.intake-title {
    text-align: center;
    font-size: 2rem;
    font-weight: 800;
    letter-spacing: -0.035em;
    color: #1d1d1f;
    margin: 18px 0 6px;
}
.intake-sub {
    text-align: center;
    color: #6e6e73;
    font-size: 1rem;
    line-height: 1.55;
    margin: 0 auto 24px;
    max-width: 520px;
}
.intake-formats {
    text-align: center;
    color: #86868b;
    font-size: 0.78rem;
    letter-spacing: 0.01em;
    margin: 10px 0 4px;
}
/* Centre + constrain the file uploader so it reads like one clean bar */
[data-testid="stFileUploader"] {
    max-width: 680px;
    margin: 0 auto 8px;
}
[data-testid="stFileUploaderDropzone"] {
    min-height: 92px;
    align-items: center;
}
/* First-question search box (a Streamlit form) */
[data-testid="stForm"] {
    max-width: 680px;
    margin: 6px auto 0;
    border: none !important;
    padding: 0 !important;
    background: transparent !important;
}
[data-testid="stForm"] [data-testid="stTextInput"] input {
    border-radius: 980px !important;
    border: 1px solid #d2d2d7 !important;
    background: #ffffff !important;
    padding: 15px 22px !important;
    font-size: 1rem !important;
    box-shadow: 0 4px 18px rgba(0,0,0,0.07) !important;
    transition: border-color 0.15s ease, box-shadow 0.15s ease !important;
}
[data-testid="stForm"] [data-testid="stTextInput"] input:focus {
    border-color: #0071e3 !important;
    box-shadow: 0 4px 22px rgba(0,113,227,0.18) !important;
    outline: none !important;
}
[data-testid="stFormSubmitButton"] > button {
    border-radius: 980px !important;
    background: #1d1d1f !important;
    color: #ffffff !important;
    border: 1px solid #1d1d1f !important;
    font-weight: 600 !important;
    font-size: 0.86rem !important;
    padding: 9px 22px !important;
    transition: background 0.15s ease, transform 0.15s ease !important;
}
[data-testid="stFormSubmitButton"] > button:hover {
    background: #000000 !important;
    transform: translateY(-1px) !important;
}

/* ── Responsive / phone layout ── */
@media (max-width: 768px) {
    .block-container { padding: 1rem 0.75rem !important; }
    .hero { padding: 22px 20px; border-radius: 18px; }
    .hero-title { font-size: 1.55rem; }
    .hero-desc { font-size: 0.86rem; }
    .welcome-title { font-size: 1.9rem; }
    .welcome-sub { font-size: 0.95rem; }
    .intake-title { font-size: 1.5rem; }
    .intake-sub { font-size: 0.9rem; }
    .feature-grid { grid-template-columns: 1fr; max-width: 100%; }
    .mobile-hint { display: block; }
    [data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-user"]) [data-testid="stMarkdownContainer"],
    [data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-assistant"]) [data-testid="stMarkdownContainer"] {
        max-width: 94% !important;
    }
}
</style>
""",
    unsafe_allow_html=True,
)


# ── Secrets ────────────────────────────────────────────────────────────────
try:
    gemini_key = st.secrets.get("GEMINI_API_KEY", "")
except Exception:
    gemini_key = ""

try:
    groq_key = st.secrets.get("GROQ_API_KEY", "")
except Exception:
    groq_key = ""


# ── Sidebar ────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("### Model")
    llm_provider = st.selectbox(
        "Provider",
        ["Gemini", "Groq"],
        help=(
            "Gemini: free tier via Google AI Studio. "
            "Groq: free API, no credit card, fast Llama 3.3 70B inference."
        ),
        label_visibility="collapsed",
    )

    st.markdown("### Retrieval")
    use_hybrid_search = st.toggle("Hybrid BM25 + FAISS", value=True)
    use_reranker = st.toggle("Cross-encoder reranking", value=True)
    use_query_expansion = st.toggle(
        "Query expansion",
        value=False,
        help="Generates extra query variants using the LLM — one extra API call per question.",
    )
    top_k_context = st.slider("Cited chunks", 2, 8, TOP_K_CONTEXT)
    show_diagnostics = st.toggle("Retrieval diagnostics", value=False)


# ── Cached model loaders ───────────────────────────────────────────────────
@st.cache_resource
def load_embedding_model():
    return SentenceTransformer(EMBEDDING_MODEL_NAME)


@st.cache_resource
def load_reranker():
    return CrossEncoder(RERANKER_MODEL_NAME)


# ── Utilities ──────────────────────────────────────────────────────────────
def ensure_storage_dirs():
    CACHE_DIR.mkdir(exist_ok=True)
    CHAT_DIR.mkdir(exist_ok=True)


def safe_filename(value):
    return re.sub(r"[^a-zA-Z0-9_.-]", "_", value)


def file_signature(files):
    digest = hashlib.sha256()
    digest.update(PARSER_VERSION.encode("utf-8"))
    for file in files:
        data = file.getvalue()
        digest.update(file.name.encode("utf-8", errors="ignore"))
        digest.update(str(len(data)).encode("utf-8"))
        digest.update(hashlib.sha256(data).digest())
    return digest.hexdigest()


def chat_history_path(signature):
    return CHAT_DIR / f"{safe_filename(signature)}.json"


def default_chat_message():
    return [{"role": "assistant", "content": "Documents indexed. Ask me anything about them."}]


def load_chat_history(signature):
    path = chat_history_path(signature)
    if not path.exists():
        return default_chat_message()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list) and all(isinstance(m, dict) and "role" in m for m in data):
            return data
        return default_chat_message()
    except Exception:
        return default_chat_message()


def save_chat_history(signature, messages):
    try:
        ensure_storage_dirs()
        chat_history_path(signature).write_text(
            json.dumps(messages, indent=2, ensure_ascii=False), encoding="utf-8"
        )
    except Exception as exc:
        st.warning(f"Could not save chat history: {exc}")


def export_chat_markdown(messages):
    lines = ["# RAG Chat Export", ""]
    for m in messages:
        role = m.get("role", "message").title()
        lines.extend([f"## {role}", "", m.get("content", ""), ""])
        if m.get("sources"):
            lines.extend(["### Sources", "", m["sources"], ""])
    return "\n".join(lines)


def _strip_markdown(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", str(text))   # bold
    text = re.sub(r"`([^`]*)`", r"\1", text)             # inline code
    text = re.sub(r"^>\s?", "", text, flags=re.MULTILINE)  # blockquote
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)  # headings
    return text


def export_chat_pdf(messages):
    """Render the conversation to a black-headed, Apple-styled PDF (bytes)."""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    NEXT_LINE = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT}

    def safe(text):
        # fpdf core fonts are latin-1; replace anything outside it.
        return _strip_markdown(text).encode("latin-1", "replace").decode("latin-1")

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    content_w = pdf.w - pdf.l_margin - pdf.r_margin

    # Black header banner
    pdf.set_fill_color(29, 29, 31)
    pdf.rect(0, 0, pdf.w, 30, style="F")
    pdf.set_xy(pdf.l_margin, 8)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 20)
    pdf.cell(0, 9, "RAG Analyzer", **NEXT_LINE)
    pdf.set_x(pdf.l_margin)
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(190, 190, 195)
    pdf.cell(0, 6, "Chat Transcript", **NEXT_LINE)

    pdf.ln(14)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(134, 134, 139)
    pdf.cell(0, 6, safe(f"Exported {time.strftime('%Y-%m-%d %H:%M')}"), **NEXT_LINE)
    pdf.ln(3)

    for m in messages:
        role = m.get("role", "message").title()
        pdf.set_font("Helvetica", "B", 11)
        if role.lower().startswith("user"):
            pdf.set_text_color(29, 29, 31)
        else:
            pdf.set_text_color(0, 113, 227)
        pdf.set_x(pdf.l_margin)
        pdf.cell(0, 7, safe(role), **NEXT_LINE)

        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(55, 55, 60)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(content_w, 5.6, safe(m.get("content", "")))
        pdf.ln(1)

        if m.get("sources"):
            pdf.set_font("Helvetica", "I", 8)
            pdf.set_text_color(134, 134, 139)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(content_w, 4.6, safe("Sources: " + m["sources"]))
        pdf.ln(5)

    return bytes(pdf.output())


def make_source_label(metadata):
    label = metadata.get("source", "Unknown source")
    details = []
    if metadata.get("page"):
        details.append(f"page {metadata['page']}")
    if metadata.get("sheet"):
        details.append(f"sheet {metadata['sheet']}")
    if metadata.get("row_range"):
        details.append(f"rows {metadata['row_range']}")
    if metadata.get("paragraph_range"):
        details.append(f"paragraphs {metadata['paragraph_range']}")
    if details:
        label += " (" + ", ".join(details) + ")"
    return label


# ── Document parsing ───────────────────────────────────────────────────────
def dataframe_to_sections(dataframe, filename, sheet_name=None, rows_per_section=40):
    sections = []
    total_rows = len(dataframe)
    if total_rows == 0:
        sections.append({
            "text": dataframe.to_string(index=False),
            "metadata": {"source": filename, "type": "spreadsheet", "sheet": sheet_name, "row_range": "empty"},
        })
        return sections
    for start in range(0, total_rows, rows_per_section):
        end = min(start + rows_per_section, total_rows)
        sections.append({
            "text": dataframe.iloc[start:end].to_string(index=False),
            "metadata": {"source": filename, "type": "spreadsheet", "sheet": sheet_name, "row_range": f"{start + 1}-{end}"},
        })
    return sections


def extract_sections_from_file(file):
    name = file.name
    lower_name = name.lower()
    sections = []
    try:
        file.seek(0)
    except Exception:
        pass
    try:
        if lower_name.endswith(".pdf"):
            reader = pypdf.PdfReader(file)
            for page_number, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text()
                except Exception:
                    continue
                if text and text.strip():
                    sections.append({"text": text, "metadata": {"source": name, "type": "pdf", "page": page_number}})
            if not sections:
                st.warning(f"No extractable text in {name}. May be a scanned PDF.")
            return sections

        if lower_name.endswith(".docx"):
            document = docx.Document(file)
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            for start in range(0, len(paragraphs), 20):
                end = min(start + 20, len(paragraphs))
                sections.append({
                    "text": "\n".join(paragraphs[start:end]),
                    "metadata": {"source": name, "type": "docx", "paragraph_range": f"{start + 1}-{end}"},
                })
            return sections

        if lower_name.endswith(".txt"):
            data = file.read()
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                try:
                    text = data.decode("latin-1")
                except UnicodeDecodeError:
                    text = data.decode("utf-8", errors="replace")
            return [{"text": text, "metadata": {"source": name, "type": "text"}}]

        if lower_name.endswith(".csv"):
            try:
                dataframe = pd.read_csv(file)
            except UnicodeDecodeError:
                file.seek(0)
                dataframe = pd.read_csv(file, encoding="latin-1")
            except pd.errors.EmptyDataError:
                st.warning(f"{name} appears to be empty.")
                return sections
            return dataframe_to_sections(dataframe, name)

        if lower_name.endswith((".xlsx", ".xls")):
            sheets = pd.read_excel(file, sheet_name=None)
            for sheet_name, dataframe in sheets.items():
                sections.extend(dataframe_to_sections(dataframe, name, sheet_name=sheet_name))
            return sections
    except Exception as exc:
        st.error(f"Error reading {file.name}: {exc}")
    return sections


# ── Chunking ───────────────────────────────────────────────────────────────
def recursive_split_text(text, metadata, chunk_size=1000, overlap=200):
    separators = ["\n\n", "\n", ". ", " ", ""]

    def split_recursive(text_to_split, separators_list):
        if len(text_to_split) <= chunk_size:
            return [text_to_split.strip()]
        separator = separators_list[0]
        next_separators = separators_list[1:]
        if not next_separators:
            step = max(1, chunk_size - overlap)
            return [text_to_split[i: i + chunk_size].strip() for i in range(0, len(text_to_split), step)]
        pieces = text_to_split.split(separator)
        chunks, current = [], ""
        for piece in pieces:
            candidate = f"{current}{separator}{piece}" if current else piece
            if len(candidate) <= chunk_size:
                current = candidate
                continue
            if current:
                chunks.append(current.strip())
            if len(piece) > chunk_size:
                chunks.extend(split_recursive(piece, next_separators))
                current = ""
            else:
                previous = chunks[-1] if chunks else ""
                prefix = previous[-overlap:] if previous else ""
                candidate = f"{prefix}{separator}{piece}" if prefix else piece
                current = candidate if len(candidate) <= chunk_size else piece
        if current:
            chunks.append(current.strip())
        return [c for c in chunks if c]

    raw_chunks = [c for c in split_recursive(text, separators) if c and c.strip()]
    source = metadata.get("source", "document")
    return [
        {
            "id": f"{source}::chunk-{i + 1}::{hashlib.sha1(c.encode('utf-8', errors='ignore')).hexdigest()[:8]}",
            "text": c,
            "source": source,
            "source_label": make_source_label(metadata),
            "chunk_number": i + 1,
            "metadata": metadata.copy(),
        }
        for i, c in enumerate(raw_chunks)
    ]


def tokenize(text):
    return re.findall(r"[a-zA-Z0-9]+", text.lower())


# ── FAISS index ────────────────────────────────────────────────────────────
def build_faiss_index(embeddings):
    normalized = embeddings.astype("float32")
    faiss.normalize_L2(normalized)
    index = faiss.IndexFlatIP(normalized.shape[1])
    index.add(normalized)
    return index


def cache_paths(signature):
    base = CACHE_DIR / signature
    return {"index": base.with_suffix(".faiss"), "metadata": base.with_suffix(".pkl")}


def remove_index_cache(signature):
    for path in cache_paths(signature).values():
        if path.exists():
            path.unlink()


def save_index_cache(signature, index, chunks, embeddings, file_stats):
    try:
        ensure_storage_dirs()
        paths = cache_paths(signature)
        faiss.write_index(index, str(paths["index"]))
        with paths["metadata"].open("wb") as handle:
            pickle.dump(
                {"chunks": chunks, "embeddings": embeddings, "file_stats": file_stats,
                 "created_at": time.time(), "parser_version": PARSER_VERSION},
                handle,
            )
    except Exception as exc:
        st.warning(f"Could not write index cache ({exc}). Will re-index next session.")


def load_index_cache(signature):
    paths = cache_paths(signature)
    if not paths["index"].exists() or not paths["metadata"].exists():
        return None
    try:
        index = faiss.read_index(str(paths["index"]))
        with paths["metadata"].open("rb") as handle:
            meta = pickle.load(handle)
        if meta.get("parser_version") != PARSER_VERSION:
            return None
        return {
            "faiss_index": index,
            "chunks": meta["chunks"],
            "embeddings": meta["embeddings"],
            "file_stats": meta["file_stats"],
            "loaded_from_cache": True,
        }
    except Exception:
        return None


def build_document_index(files, model, signature):
    cached = load_index_cache(signature)
    if cached:
        chunks = cached["chunks"]
        cached["bm25"] = BM25Okapi([tokenize(c["text"]) for c in chunks])
        return cached

    chunks, file_stats = [], []
    for file in files:
        sections = extract_sections_from_file(file)
        file_chunks = []
        for section in sections:
            if section["text"].strip():
                file_chunks.extend(recursive_split_text(section["text"], section["metadata"]))

        truncated = False
        if len(file_chunks) > MAX_CHUNKS_PER_FILE:
            file_chunks = file_chunks[:MAX_CHUNKS_PER_FILE]
            truncated = True
            st.warning(f"{file.name} exceeded {MAX_CHUNKS_PER_FILE} chunks — first portion indexed only.")

        chunks.extend(file_chunks)
        if file_chunks:
            file_stats.append({
                "name": file.name,
                "size": f"{file.size / 1024:.1f} KB",
                "chunks": len(file_chunks),
                "sections": len(sections),
                "truncated": truncated,
            })

    if not chunks:
        return None

    try:
        embeddings = model.encode([c["text"] for c in chunks], show_progress_bar=False)
    except Exception as exc:
        st.error(f"Failed to generate embeddings: {exc}")
        st.stop()

    embeddings = np.asarray(embeddings, dtype="float32")
    index = build_faiss_index(embeddings)
    bm25 = BM25Okapi([tokenize(c["text"]) for c in chunks])
    save_index_cache(signature, index, chunks, embeddings, file_stats)

    return {
        "faiss_index": index,
        "chunks": chunks,
        "embeddings": embeddings,
        "file_stats": file_stats,
        "bm25": bm25,
        "loaded_from_cache": False,
    }


# ── Query expansion ────────────────────────────────────────────────────────
def _expansion_prompt(query):
    return (
        "Rewrite this search query into three short retrieval queries.\n"
        "Keep the original meaning. Return one query per line, no bullets or numbers.\n\n"
        f"Query: {query}"
    )


def _parse_expansions(text, original):
    expansions = [line.strip("-* 0123456789.").strip() for line in text.splitlines() if line.strip()]
    return [original] + expansions[:3]


def _groq_json_request(prompt, api_token, timeout=30):
    payload = json.dumps({
        "model": GROQ_MODEL_NAME,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.2,
        "max_tokens": 256,
    }).encode("utf-8")
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {api_token}",
        "User-Agent": HTTP_USER_AGENT,
    }
    req = urllib.request.Request(GROQ_ENDPOINT, data=payload, headers=headers)
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read())


def expand_query(query, provider, gemini_ready=False, groq_key=None):
    if provider == "Gemini" and gemini_ready:
        try:
            response = genai.GenerativeModel(GEMINI_MODEL_NAME).generate_content(_expansion_prompt(query))
            return _parse_expansions(response.text, query)
        except Exception:
            return [query]

    if provider == "Groq" and groq_key:
        try:
            data = _groq_json_request(_expansion_prompt(query), groq_key)
            return _parse_expansions(data["choices"][0]["message"]["content"], query)
        except Exception:
            return [query]

    return [query]


# ── Retrieval ──────────────────────────────────────────────────────────────
def retrieve_candidates(query, model, document_index, use_hybrid, use_expansion,
                        provider=None, gemini_ready=False, groq_key=None):
    chunks = document_index["chunks"]
    queries = (
        expand_query(query, provider, gemini_ready, groq_key) if use_expansion else [query]
    )
    candidate_scores = {}

    for expanded_query in queries:
        query_embedding = np.asarray([model.encode(expanded_query)], dtype="float32")
        faiss.normalize_L2(query_embedding)
        distances, indices = document_index["faiss_index"].search(
            query_embedding, min(TOP_K_RETRIEVAL, len(chunks))
        )
        for score, index in zip(distances[0], indices[0]):
            if index == -1:
                continue
            d = candidate_scores.setdefault(index, {"index": index, "semantic": 0.0, "bm25": 0.0, "combined": 0.0})
            d["semantic"] = max(d["semantic"], float(score))

    if use_hybrid:
        bm25 = document_index.get("bm25") or BM25Okapi([tokenize(c["text"]) for c in chunks])
        bm25_scores = bm25.get_scores(tokenize(query))
        if np.max(bm25_scores) > 0:
            bm25_scores = bm25_scores / np.max(bm25_scores)
        for index in np.argsort(bm25_scores)[-TOP_K_RETRIEVAL:][::-1]:
            d = candidate_scores.setdefault(int(index), {"index": int(index), "semantic": 0.0, "bm25": 0.0, "combined": 0.0})
            d["bm25"] = max(d["bm25"], float(bm25_scores[index]))

    for d in candidate_scores.values():
        d["combined"] = d["semantic"] + (0.35 * d["bm25"])

    return sorted(candidate_scores.values(), key=lambda x: x["combined"], reverse=True)[:TOP_K_RETRIEVAL]


def rerank_candidates(query, candidates, chunks):
    if not candidates:
        return []
    try:
        reranker = load_reranker()
        pairs = [(query, chunks[c["index"]]["text"]) for c in candidates]
        scores = reranker.predict(pairs)
        for c, score in zip(candidates, scores):
            c["rerank"] = float(score)
        return sorted(candidates, key=lambda x: x.get("rerank", x["combined"]), reverse=True)
    except Exception as exc:
        st.warning(f"Reranking failed, using retrieval scores: {exc}")
        return sorted(candidates, key=lambda x: x["combined"], reverse=True)


def build_context(selected_candidates, chunks):
    context_parts, source_parts, diagnostics = [], [], []
    for rank, candidate in enumerate(selected_candidates, start=1):
        chunk = chunks[candidate["index"]]
        citation_id = f"{chunk['source_label']} - chunk {chunk['chunk_number']}"
        context_parts.append(f"[{rank}] Source: {citation_id}\n{chunk['text']}")
        snippet = chunk["text"].strip()
        if len(snippet) > 900:
            snippet = f"{snippet[:900]}..."
        score_text = (
            f"semantic={candidate.get('semantic', 0):.3f}, "
            f"bm25={candidate.get('bm25', 0):.3f}, "
            f"combined={candidate.get('combined', 0):.3f}"
        )
        if "rerank" in candidate:
            score_text += f", rerank={candidate['rerank']:.3f}"
        source_parts.append(f"**[{rank}] {citation_id}**\n\nScore: {score_text}\n\n> {snippet}")
        diagnostics.append({
            "rank": rank,
            "source": citation_id,
            "semantic": round(candidate.get("semantic", 0), 4),
            "bm25": round(candidate.get("bm25", 0), 4),
            "combined": round(candidate.get("combined", 0), 4),
            "rerank": round(candidate["rerank"], 4) if "rerank" in candidate else None,
        })
    return "\n\n---\n\n".join(context_parts), "\n\n---\n\n".join(source_parts), diagnostics


def recent_conversation(messages, max_messages=6):
    return "\n".join(
        f"{m.get('role', 'user').title()}: {m.get('content', '')}"
        for m in messages[-max_messages:]
    )


# ── LLM streaming ──────────────────────────────────────────────────────────
def stream_gemini_response(prompt):
    model = genai.GenerativeModel(GEMINI_MODEL_NAME)
    for chunk in model.generate_content(prompt, stream=True, request_options={"timeout": 120}):
        text = getattr(chunk, "text", "")
        if text:
            yield text


def stream_groq_response(prompt, api_token):
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {api_token}",
        "User-Agent": HTTP_USER_AGENT,
    }
    payload = json.dumps({
        "model": GROQ_MODEL_NAME,
        "messages": [{"role": "user", "content": prompt}],
        "stream": True,
        "temperature": 0.2,
        "max_tokens": 2048,
    }).encode("utf-8")
    req = urllib.request.Request(GROQ_ENDPOINT, data=payload, headers=headers)
    try:
        resp = urllib.request.urlopen(req, timeout=120)
    except urllib.error.HTTPError as exc:
        # Groq explains *why* it rejected us in the body — surface it.
        try:
            body = exc.read().decode("utf-8", errors="replace")
        except Exception:
            body = ""
        raise RuntimeError(f"HTTP {exc.code}: {body or exc.reason}") from exc
    with resp:
        for line in resp:
            line_str = line.decode("utf-8").strip()
            if not line_str or not line_str.startswith("data: "):
                continue
            data_content = line_str[6:]
            if data_content == "[DONE]":
                break
            try:
                data = json.loads(data_content)
                content = data.get("choices", [{}])[0].get("delta", {}).get("content", "")
                if content:
                    yield content
            except Exception:
                continue


def stream_llm_response(prompt, provider, **kwargs):
    if provider == "Gemini":
        yield from stream_gemini_response(prompt)
    elif provider == "Groq":
        yield from stream_groq_response(prompt, kwargs["groq_key"])


def friendly_llm_error(error, provider):
    msg = str(error)
    low = msg.lower()
    if provider == "Gemini":
        if "429" in msg or "quota" in low or "rate" in low:
            m = re.search(r"retry[_ ]delay\s*\{\s*seconds:\s*(\d+)", msg, re.IGNORECASE)
            secs = m.group(1) if m else "30"
            return (
                f"Gemini rate limit hit (free tier). Wait ~{secs}s then retry. "
                "Tip: turn off Query expansion to reduce API calls. "
                "Or switch to Groq in the sidebar."
            )
        return f"Gemini error: {msg}"
    if provider == "Groq":
        if "401" in msg or "invalid_api_key" in low:
            return "Groq auth failed — check your key at console.groq.com/keys."
        if "429" in msg or "rate" in low:
            return "Groq rate limit hit (free tier). Wait a few seconds and retry, or switch to Gemini."
        if "413" in msg or "too large" in low or "context" in low:
            return "Request too large for the model context. Reduce cited chunks in the sidebar and retry."
        return f"Groq error: {msg}"
    return f"Error: {msg}"


def build_prompt(messages, context, query):
    return f"""You are a careful RAG assistant.
Answer the user's question using only the provided document context.
Use concise, professional language.
Add inline citations like [1] or [2] when you use evidence.
If the answer is not supported by the context, say the documents do not contain enough information.

Recent conversation:
{recent_conversation(messages)}

Document context:
{context}

Question:
{query}

Answer:"""


# ── Upload validation ──────────────────────────────────────────────────────
def validate_uploads(files):
    oversized = [f.name for f in files if f.size > MAX_FILE_SIZE_MB * 1024 * 1024]
    total_mb = sum(f.size for f in files) / (1024 * 1024)
    errors = []
    if oversized:
        errors.append(f"Files over {MAX_FILE_SIZE_MB} MB limit (skipped): " + ", ".join(oversized))
    if total_mb > MAX_TOTAL_UPLOAD_MB:
        errors.append(f"Total upload {total_mb:.1f} MB exceeds {MAX_TOTAL_UPLOAD_MB} MB limit.")
    accepted = [f for f in files if f.size <= MAX_FILE_SIZE_MB * 1024 * 1024]
    return accepted, errors


# ── State ──────────────────────────────────────────────────────────────────
gemini_ready = llm_provider == "Gemini" and bool(gemini_key)
groq_ready = llm_provider == "Groq" and bool(groq_key)
api_ready = gemini_ready or groq_ready

if gemini_ready:
    genai.configure(api_key=gemini_key)


# ── Main area ──────────────────────────────────────────────────────────────
# Always-visible hero — describes the app and is the first thing seen on phones,
# where the sidebar is collapsed behind the menu button.
st.markdown(
    """
    <div class="hero">
        <div class="hero-badge">Retrieval-Augmented Generation</div>
        <div class="hero-title">RAG Analyzer</div>
        <div class="hero-desc">
            Upload your PDFs, Word docs, spreadsheets, or text files and ask questions in plain English.
            The app splits each document into passages, finds the ones that actually answer your
            question using <strong>hybrid semantic + keyword search</strong>, and generates a grounded,
            <strong>cited</strong> answer with a free LLM — no data leaves for embeddings, and it costs nothing to run.
        </div>
    </div>
    <div class="mobile-hint">
        📱 <strong>On a phone?</strong> Tap the <strong>›</strong> arrow at the top-left to open the menu —
        that's where you pick your model and tune retrieval settings.
    </div>
    """,
    unsafe_allow_html=True,
)

if not api_ready:
    st.markdown(
        """
        <div class="welcome-wrap">
            <div class="welcome-icon">🔧</div>
            <div class="welcome-title">Almost ready</div>
            <div class="welcome-sub">
                This deployment is being configured. The AI models will be available shortly —
                no action needed on your end.
            </div>
            <div class="feature-grid">
                <div class="feature-card">
                    <div class="feature-card-title">Hybrid Retrieval</div>
                    <div class="feature-card-desc">FAISS + BM25 with cross-encoder reranking</div>
                </div>
                <div class="feature-card">
                    <div class="feature-card-title">Zero Cost</div>
                    <div class="feature-card-desc">Local embeddings · Free Gemini or Groq</div>
                </div>
                <div class="feature-card">
                    <div class="feature-card-title">Multi-format</div>
                    <div class="feature-card-desc">PDF, Word, TXT, CSV, and Excel</div>
                </div>
                <div class="feature-card">
                    <div class="feature-card-title">Cited Answers</div>
                    <div class="feature-card-desc">Page-aware citations from every chunk</div>
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

else:
    # ── Central intake — the "search bar" ────────────────────────────────
    prior = st.session_state.get("messages", [])
    started = any(m.get("role") == "user" for m in prior)

    uploader_kwargs = dict(
        label="Attach documents",
        type=["pdf", "docx", "txt", "csv", "xlsx", "xls"],
        accept_multiple_files=True,
        key="uploader",
        label_visibility="collapsed",
    )

    if started:
        # Once chatting, keep the uploader tucked away so the thread stays clean.
        with st.expander("📎  Documents", expanded=False):
            uploaded_files = st.file_uploader(**uploader_kwargs)
    else:
        st.markdown(
            '<div class="intake-title">Chat with your documents</div>'
            '<div class="intake-sub">Attach a file, then ask anything — '
            'every answer is grounded in your document and cited.</div>',
            unsafe_allow_html=True,
        )
        uploaded_files = st.file_uploader(**uploader_kwargs)

    if not uploaded_files:
        st.markdown(
            f'<div class="intake-formats">PDF · Word · Excel · CSV · TXT'
            f' &nbsp;·&nbsp; up to {MAX_FILE_SIZE_MB} MB per file</div>'
            '<div class="feature-grid" style="margin: 32px auto 0;">'
            '<div class="feature-card"><div class="feature-card-title">Hybrid Retrieval</div>'
            '<div class="feature-card-desc">FAISS + BM25 with cross-encoder reranking</div></div>'
            '<div class="feature-card"><div class="feature-card-title">Cited Answers</div>'
            '<div class="feature-card-desc">Page-aware citations from every chunk</div></div>'
            '<div class="feature-card"><div class="feature-card-title">Multi-format</div>'
            '<div class="feature-card-desc">PDF, Word, TXT, CSV, and Excel</div></div>'
            '<div class="feature-card"><div class="feature-card-title">Zero Cost</div>'
            '<div class="feature-card-desc">Local embeddings · Free Gemini or Groq</div></div>'
            '</div>',
            unsafe_allow_html=True,
        )
        st.stop()

    # ── Validate & index ─────────────────────────────────────────────────
    uploaded_files, upload_errors = validate_uploads(uploaded_files)
    for err in upload_errors:
        st.error(err)
    if not uploaded_files:
        st.stop()

    ensure_storage_dirs()
    signature = file_signature(uploaded_files)

    with st.sidebar:
        st.markdown("### Session")
        col_a, col_b = st.columns(2)
        with col_a:
            if st.button("Clear chat", use_container_width=True):
                st.session_state["messages"] = default_chat_message()
                save_chat_history(signature, st.session_state["messages"])
                st.rerun()
        with col_b:
            if st.button("Rebuild index", use_container_width=True):
                remove_index_cache(signature)
                st.session_state.pop("file_signature", None)
                st.session_state.pop("document_index", None)
                st.rerun()

    with st.spinner("Loading embedding model…"):
        embedding_model = load_embedding_model()

    if st.session_state.get("file_signature") != signature:
        with st.spinner("Indexing documents…"):
            document_index = build_document_index(uploaded_files, embedding_model, signature)
            if not document_index:
                st.error("No extractable text found in the uploaded files.")
                st.stop()
            st.session_state["document_index"] = document_index
            st.session_state["file_signature"] = signature
            st.session_state["messages"] = load_chat_history(signature)

    document_index = st.session_state["document_index"]
    chunks = document_index["chunks"]
    if "messages" not in st.session_state:
        st.session_state["messages"] = load_chat_history(signature)
    messages = st.session_state["messages"]
    started = any(m.get("role") == "user" for m in messages)

    # ── Metrics ──────────────────────────────────────────────────────────
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("Documents", len(uploaded_files))
    c2.metric("Chunks", len(chunks))
    c3.metric("Search", "Hybrid" if use_hybrid_search else "FAISS")
    c4.metric("Index", "Cached" if document_index["loaded_from_cache"] else "Fresh")
    c5.metric("Model", llm_provider)

    with st.expander("Indexed files", expanded=False):
        for stat in document_index.get("file_stats", []):
            safe_name = html.escape(str(stat["name"]))
            truncated_note = ' <span style="color:#b45309;">(truncated)</span>' if stat.get("truncated") else ""
            st.markdown(
                f'<div class="doc-card"><strong>{safe_name}</strong>{truncated_note}'
                f" &nbsp;·&nbsp; {stat['size']} &nbsp;·&nbsp; {stat['chunks']} chunks</div>",
                unsafe_allow_html=True,
            )

    try:
        st.download_button(
            "⬇  Download transcript (PDF)",
            data=export_chat_pdf(messages),
            file_name="rag_chat_transcript.pdf",
            mime="application/pdf",
        )
    except Exception:
        st.download_button(
            "⬇  Download transcript (Markdown)",
            data=export_chat_markdown(messages),
            file_name="rag_chat_transcript.md",
            mime="text/markdown",
        )

    # ── First-question composer (the centred search box) ─────────────────
    if not started:
        st.markdown(
            '<div class="intake-sub" style="margin: 24px auto 12px;">'
            'Your document is indexed. Ask your first question below.</div>',
            unsafe_allow_html=True,
        )
        with st.form("first_query", clear_on_submit=True):
            first_q = st.text_input(
                "First question",
                label_visibility="collapsed",
                placeholder="Ask a question about your document…",
            )
            asked = st.form_submit_button("Ask", use_container_width=True)
        if asked and first_q.strip():
            messages.append({"role": "user", "content": first_q.strip()})
            save_chat_history(signature, messages)
            st.rerun()

    # ── Conversation thread ──────────────────────────────────────────────
    for message in messages:
        with st.chat_message(message["role"]):
            st.write(message["content"])
            if message.get("sources"):
                with st.expander("Sources", expanded=False):
                    st.markdown(message["sources"])
            if show_diagnostics and message.get("diagnostics"):
                with st.expander("Retrieval diagnostics", expanded=False):
                    st.dataframe(pd.DataFrame(message["diagnostics"]), use_container_width=True)

    # ── Generate the answer when the last turn is still the user's ───────
    if messages and messages[-1]["role"] == "user":
        user_query = messages[-1]["content"]
        with st.spinner("Retrieving evidence…"):
            relevant_context, sources_string, diagnostics = "", "", []
            try:
                candidates = retrieve_candidates(
                    user_query,
                    embedding_model,
                    document_index,
                    use_hybrid_search,
                    use_query_expansion,
                    provider=llm_provider,
                    gemini_ready=gemini_ready,
                    groq_key=groq_key if groq_ready else None,
                )
                if use_reranker:
                    candidates = rerank_candidates(user_query, candidates, chunks)
                selected_candidates = candidates[:top_k_context]
                relevant_context, sources_string, diagnostics = build_context(selected_candidates, chunks)
            except Exception as exc:
                st.error(f"Retrieval failed: {exc}")

        prompt = build_prompt(messages, relevant_context, user_query)

        with st.chat_message("assistant"):
            placeholder = st.empty()
            streamed_text = ""
            error_note = None
            llm_kwargs = {"groq_key": groq_key} if llm_provider == "Groq" else {}

            for attempt in range(LLM_MAX_RETRIES + 1):
                streamed_text = ""
                try:
                    for text_chunk in stream_llm_response(prompt, llm_provider, **llm_kwargs):
                        streamed_text += text_chunk
                        placeholder.markdown(streamed_text)
                    error_note = None
                    break
                except Exception as exc:
                    error_note = exc
                    is_transient = "429" not in str(exc) and "quota" not in str(exc).lower()
                    if attempt < LLM_MAX_RETRIES and is_transient:
                        time.sleep(LLM_RETRY_BASE_DELAY_SECONDS * (attempt + 1))
                        continue
                    break

            if error_note is not None:
                err_msg = friendly_llm_error(error_note, llm_provider)
                streamed_text = (
                    f"{streamed_text}\n\n---\n*Response cut off: {err_msg}*"
                    if streamed_text.strip()
                    else err_msg
                )
                placeholder.markdown(streamed_text)

            if sources_string:
                with st.expander("Sources", expanded=False):
                    st.markdown(sources_string)
            if show_diagnostics and diagnostics:
                with st.expander("Retrieval diagnostics", expanded=False):
                    st.dataframe(pd.DataFrame(diagnostics), use_container_width=True)

        messages.append({
            "role": "assistant",
            "content": streamed_text,
            "sources": sources_string,
            "diagnostics": diagnostics,
        })
        save_chat_history(signature, messages)

    # ── Follow-up input, pinned to the bottom once the chat has started ──
    if started:
        if user_query := st.chat_input("Ask a follow-up…"):
            messages.append({"role": "user", "content": user_query})
            save_chat_history(signature, messages)
            st.rerun()
